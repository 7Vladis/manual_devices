"""Выгрузка справочника в XLSX — опись оборудования для сверки на месте."""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.utils import timezone

from ..models import DataObject
from users.decorators import role_required


# Excel/LibreOffice трактуют значение, начинающееся с =, +, -, @ (а также
# с управляющих символов табуляции и возврата каретки) как формулу.
# Данные приходят от пользователей и из YouTrack, поэтому обезвреживаем их.
FORMULA_PREFIXES = ('=', '+', '-', '@', '\t', '\r')

# Серым помечаются объекты без инвентарного номера: на бумаге их видно сразу,
# а из описи они не исчезают — иначе вместе с ними разваливалось бы дерево.
# Узел без номера держит под собой пронумерованные, и выбросить его нельзя.
NO_INVENTORY_FILL = 'F1F5F9'
INK = '1E293B'
RULE = 'CBD5E1'


def escape_formula(value):
    text = "" if value is None else str(value)
    if text.startswith(FORMULA_PREFIXES):
        return "'" + text
    return text


def walk_tree(objects):
    """
    Обход дерева сверху вниз: родитель, под ним его состав, и так вглубь.

    Порядок строк — это маршрут обхода помещения: пришёл к агрегату, и рядом
    на листе весь его состав. Прежняя сортировка по инвентарному номеру
    разносила узлы одного агрегата по всей описи, потому что у компонентов
    номера обычно нет вовсе.

    Возвращает тройки «объект, структурный номер, уровень». Структурный номер
    (1, 1.2, 1.2.1) нужен для сверки вслух: «пункт один-два-один».

    Дерево собирается из одной выборки, а не запросом на узел: объектов
    немного, зато выгрузка не упирается в число обращений к базе.
    """
    children = {}
    for obj in objects:
        children.setdefault(obj.parent_id, []).append(obj)

    for siblings in children.values():
        siblings.sort(key=lambda item: (item.name or item.model.name or '').lower())

    rows = []

    def descend(parent_id, prefix, level):
        for position, obj in enumerate(children.get(parent_id, []), start=1):
            number = f"{prefix}{position}"
            rows.append((obj, number, level))
            descend(obj.uuid, f"{number}.", level + 1)

    descend(None, '', 0)
    return rows


@login_required
@role_required(['admin', 'superuser'])
def export_xlsx_view(request):
    """
    Опись оборудования в XLSX.

    Файл печатают и носят в руках: по нему комиссия ищет объект на месте
    и отмечает ручкой, что нашла. Поэтому лист настроен на печать целиком —
    альбомная ориентация, вписывание по ширине, шапка таблицы на каждой
    странице, — а всё, что мешает читать опись (описания, характеристики),
    вынесено на отдельные листы.

    Выгрузка одна и всегда полная: выбор «только с инвентарным номером»
    убран. Он рвал дерево — объект без номера исчезал вместе со своим
    составом, хотя именно он и говорит, где этот состав искать.
    """
    objects = DataObject.objects.select_related('model', 'model__object_type', 'parent')
    rows = walk_tree(list(objects))

    wb = openpyxl.Workbook()
    _fill_inventory_sheet(wb, rows)
    _fill_models_sheet(wb, rows)
    _fill_descriptions_sheet(wb, rows)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"opis_oborudovaniya_{timezone.localdate().strftime('%Y%m%d')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def _fill_inventory_sheet(wb, rows):
    """Лист «Опись» — тот, который печатают."""
    ws = wb.active
    ws.title = "Опись"

    # Шапка документа: комиссия спрашивает, на какой момент данные.
    ws['A1'] = "Опись оборудования"
    ws['A1'].font = Font(bold=True, size=14, color=INK)
    ws['A2'] = f"Выгружено {timezone.localtime().strftime('%d.%m.%Y в %H:%M')}"
    ws['A3'] = f"Всего объектов: {len(rows)}"
    for address in ('A2', 'A3'):
        ws[address].font = Font(size=10, color="64748B")

    headers = [
        "№",
        "Инвентарный номер",
        "Наименование",
        "Модель",
        "Входит в состав",
        "Следующее ТО",
        "Отметка",
        "Примечание",
    ]
    # Строка 4 остаётся пустой: шапка документа отделена от таблицы.
    header_row = 5
    for column, title in enumerate(headers, start=1):
        ws.cell(row=header_row, column=column, value=title)

    thin = Side(style='thin', color=RULE)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Шапка без заливки: сплошная тёмная полоса съедает тонер, а на печати
    # её роль и так выполняют полужирный шрифт и рамка.
    for column in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=column)
        cell.font = Font(bold=True, color=INK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    grey = PatternFill(start_color=NO_INVENTORY_FILL, end_color=NO_INVENTORY_FILL,
                       fill_type="solid")

    for obj, number, level in rows:
        parent = obj.parent
        values = [
            number,
            obj.inventory_number or "",
            obj.name or obj.model.name,
            obj.model.name,
            (parent.name or parent.model.name) if parent else "",
            obj.next_maintenance_date.strftime('%d.%m.%Y') if obj.next_maintenance_date else "",
            "",  # отметка о наличии — ставится ручкой
            "",  # примечание — пишется ручкой
        ]
        ws.append([escape_formula(value) for value in values])

        row_index = ws.max_row
        for cell in ws[row_index]:
            cell.data_type = 's'
            cell.border = border
            cell.alignment = Alignment(vertical="center")
            if not obj.inventory_number:
                cell.fill = grey

        # Уровень показан родным отступом Excel, а не пробелами в тексте:
        # он печатается, не мешает поиску и не ломает фильтр.
        ws.cell(row=row_index, column=3).alignment = Alignment(vertical="center", indent=level)

        # Группировка: на экране состав агрегата можно свернуть. На печать
        # не влияет, поэтому бланку не мешает.
        if level:
            ws.row_dimensions[row_index].outline_level = min(level, 7)

    ws.auto_filter.ref = f"A{header_row}:H{ws.max_row}"
    ws.freeze_panes = f"A{header_row + 1}"

    widths = [8, 20, 42, 24, 28, 14, 10, 28]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    # Печать. Без этого лист расползался на несколько страниц в ширину,
    # и со второй страницы уже не было видно, где какой столбец.
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.print_options.horizontalCentered = True
    ws.oddFooter.left.text = "Опись оборудования"
    ws.oddFooter.right.text = "Стр. &P из &N"


def _fill_models_sheet(wb, rows):
    """
    Характеристики — по одному разу на модель, а не на каждый экземпляр.

    Они принадлежат модели, и прежняя выгрузка повторяла их в строке каждого
    объекта: на три десятка моделей приходились сотни одинаковых длинных
    строк, которые к тому же и были самым широким столбцом описи.
    """
    ws = wb.create_sheet("Модели")
    ws.append(["Модель", "Тип оборудования", "Характеристики", "Объектов в описи"])

    for column in range(1, 5):
        cell = ws.cell(row=1, column=column)
        cell.font = Font(bold=True, color=INK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style='thin', color=RULE))

    counts = {}
    models = {}
    for obj, _number, _level in rows:
        counts[obj.model_id] = counts.get(obj.model_id, 0) + 1
        models[obj.model_id] = obj.model

    for model_id, model in sorted(models.items(), key=lambda item: item[1].name.lower()):
        specs = model.specifications if isinstance(model.specifications, dict) else {}
        specs_text = "; ".join(f"{key}: {value}" for key, value in specs.items())
        values = [model.name, model.object_type.type, specs_text, str(counts[model_id])]
        ws.append([escape_formula(value) for value in values])
        for cell in ws[ws.max_row]:
            cell.data_type = 's'
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for index, width in enumerate([32, 24, 70, 18], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"


def _fill_descriptions_sheet(wb, rows):
    """
    Описания отдельным листом.

    В описи им места нет: это самый длинный текст, он и разносил лист по
    ширине. Но выбрасывать его нельзя — в описании нередко записано точное
    место, где стоит объект, а ради этого опись и печатают.
    """
    ws = wb.create_sheet("Описания")
    ws.append(["№", "Инвентарный номер", "Наименование", "Описание"])

    for column in range(1, 5):
        cell = ws.cell(row=1, column=column)
        cell.font = Font(bold=True, color=INK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style='thin', color=RULE))

    for obj, number, _level in rows:
        if not (obj.description or '').strip():
            continue
        values = [number, obj.inventory_number or "", obj.name or obj.model.name,
                  obj.description.strip()]
        ws.append([escape_formula(value) for value in values])
        for cell in ws[ws.max_row]:
            cell.data_type = 's'
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for index, width in enumerate([8, 20, 36, 90], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
